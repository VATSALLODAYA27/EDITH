"""Turn workspace files into JSON the UI can draw: slides, document blocks, sheet tables, or plain text.

Only CONTENT and structure (not fonts/themes); the real file is always available via download.
"""
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from agents.document import _safe_path
from agents.ppt import slide_text
from userdata import workspace_dir

MAX_ROWS, MAX_COLS, MAX_TEXT = 30, 12, 5_000


def is_user_file(name: str) -> bool:
    """Skip Office lock files ("~$deck.pptx", created while a file is open in PowerPoint/Word) and hidden files."""
    return not name.startswith(("~$", "."))


def snapshot() -> dict[str, float]:
    """name -> last-modified time of every workspace file (to spot what a run created or changed)."""
    return {p.name: p.stat().st_mtime for p in workspace_dir().iterdir() if p.is_file() and is_user_file(p.name)}


def changed_since(before: dict[str, float]) -> list[str]:
    return sorted(name for name, mtime in snapshot().items() if before.get(name) != mtime)


def preview(filename: str) -> dict:
    path = _safe_path(filename)  # SECURITY: plain names inside workspace/ only
    if not path.exists():
        raise FileNotFoundError(filename)
    kind = path.suffix.lower().lstrip(".")

    if kind == "pptx":
        slides = []
        for n, s in enumerate(Presentation(path).slides, start=1):
            title, lines = slide_text(s)  # designed slides keep text in named shapes, charts become text
            slides.append({
                "number": n,
                "title": title,
                "bullets": lines,
                "notes": s.notes_slide.notes_text_frame.text if s.has_notes_slide else "",
            })
        return {"type": "pptx", "slides": slides}

    if kind == "docx":
        blocks = []
        for p in Document(path).paragraphs:
            if not p.text.strip():
                continue
            style = p.style.name.lower()
            k = "h1" if style in ("heading 1", "title") else "h2" if style.startswith("heading") else \
                "bullet" if "list" in style else "p"
            blocks.append({"kind": k, "text": p.text})
        return {"type": "docx", "blocks": blocks}

    if kind == "xlsx":
        sheets = []
        for ws in load_workbook(path).worksheets:  # formulas come back as text like "=SUM(B2:B5)"
            # ws.max_row counts "touched" but empty rows too, so keep only rows that contain something
            filled = [["" if v is None else str(v) for v in row[:MAX_COLS]]
                      for row in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in row)]
            sheets.append({"name": ws.title, "rows": filled[:MAX_ROWS], "total_rows": len(filled),
                           "charts": len(ws._charts)})
        return {"type": "xlsx", "sheets": sheets}

    if kind == "pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    elif kind in ("txt", "md"):
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"No preview for .{kind} files")
    return {"type": "text", "text": text[:MAX_TEXT]}
