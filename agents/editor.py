"""Document Editor: changes the CONTENT of existing documents, keeping their look.

    task -> tool-calling loop (list / read / fill_template / edit_document) -> "Created X" -> result

  fill_template - "rewrite this PDF in our template / on our letterhead": the template .docx is copied with its logo,
                  header, footer, fonts and page setup, and the new content goes below it in the template's styles
  edit_document - surgical find -> replace inside a .docx (body, tables, headers, footers), keeping formatting
Both save a NEW file (name_2.docx...), so the user's originals are never lost.
A separate agent from the Document agent so routing is precise: "change this document" -> here, nowhere else.

Run the standalone test from the project root:  python -m agents.editor
"""
from pathlib import Path

from docx import Document
from langchain_core.tools import tool
from pydantic import BaseModel, Field

from agents.document import _safe_path, _write_markdown, free_path, list_files, read_document
from agents.tool_agent import run_tool_agent


def _docx(filename: str) -> Path:
    path = _safe_path(filename)
    if path.suffix.lower() != ".docx" or not path.exists():
        raise FileNotFoundError(f"'{filename}' is not an existing .docx file in the workspace. Use list_files. "
                                "(A PDF can't be edited in place: read it, then write the new version with "
                                "fill_template or the Document agent.)")
    return path


def _output(filename: str, fallback: str) -> Path:
    path = _safe_path(filename or fallback)
    if path.suffix.lower() != ".docx":
        raise ValueError("output_filename must end with .docx")
    return free_path(path)  # never overwrite: a taken name becomes name_2.docx


@tool
def fill_template(template_file: str, content: str, output_filename: str = "") -> str:
    """Write content into the user's template .docx (letterhead, company template, form): the template is copied with
    its logo, header, footer, fonts and page setup, and `content` goes below what it already contains. `content` is
    simple markdown: '# ' title, '## ' headings, '- ' bullets, other lines are paragraphs. Use it for "rewrite / redo
    this document in our template". Saves a NEW file (default: <template>_filled.docx) and returns its name."""
    template = _docx(template_file)
    doc = Document(template)
    body = doc.element.body
    for p in reversed(doc.paragraphs):  # drop trailing empty lines, so content starts right below the template's
        if p.text.strip() or p._p.xpath(".//w:drawing | .//w:pict"):  # keep text and images (a logo in the body)
            break
        body.remove(p._p)
    _write_markdown(doc, content)  # appended before the section settings, so headers/footers/margins stay
    out = _output(output_filename, f"{template.stem}_filled.docx")
    doc.save(out)
    return f"Created {out.name} from the template {template_file}"


class Replacement(BaseModel):
    find: str = Field(description="Exact text currently in the document (copy it from read_document).")
    replace: str = Field(description="The new text. Empty = delete the text.")


def _all_paragraphs(doc):
    """Body, table cells (also nested), headers and footers: everywhere text can live in a .docx."""
    def walk(container):
        yield from container.paragraphs
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            yield from walk(part)


def _replace_in(paragraph, find: str, new: str) -> int:
    """Replace inside one paragraph. Inside a single run: formatting untouched. Spanning runs: the text moves into
    the first run, so the paragraph keeps its style and the first run's formatting."""
    count = paragraph.text.count(find)
    if not count:
        return 0
    for run in paragraph.runs:
        if find in run.text:
            run.text = run.text.replace(find, new)
    if find in paragraph.text:  # still there = it spans several runs
        text = paragraph.text.replace(find, new)
        runs = paragraph.runs
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    return count


@tool
def edit_document(filename: str, replacements: list[Replacement], output_filename: str = "") -> str:
    """Surgical edits to an existing .docx: replace exact text (body, tables, headers, footers) and keep all
    formatting, logo and layout. Read the document first and copy each `find` exactly. If any `find` isn't in the
    document, nothing is saved. Saves a NEW file (default: <name>_edited.docx) and returns its name."""
    source = _docx(filename)
    doc = Document(source)
    paragraphs = list(_all_paragraphs(doc))
    missing = [r.find for r in replacements if not any(r.find in p.text for p in paragraphs)]
    if missing:  # all-or-nothing: a half-applied edit is worse than none
        raise ValueError(f"Not found in {filename} (nothing changed): {missing[:3]}. Copy the exact text from "
                         "read_document; one `find` must not cross a paragraph break.")
    total = sum(_replace_in(p, r.find, r.replace) for r in replacements for p in paragraphs)
    out = _output(output_filename, f"{source.stem}_edited.docx")
    doc.save(out)
    return f"Created {out.name} with {total} replacement(s) from {filename}"


TOOLS = [list_files, read_document, fill_template, edit_document]

SYSTEM = """You are the Document Editor. You change the content of the user's documents and keep their look.
- If unsure of a file name, call list_files first. Always read_document the source before changing anything.
- "Rewrite / put this in our template" (e.g. a PDF + a template or letterhead .docx): write the full new content
  (apply the requested changes, keep every fact, number, name and date exact), then call fill_template with the
  template file. Don't repeat the template's own text (company name, address, logo text) in the content.
- Small, specific changes to a .docx (a name, a date, a sentence): edit_document with exact find/replace pairs copied
  from read_document. Prefer this over rewriting, so all formatting stays.
- PDFs can't be edited in place: write the new version with fill_template (or ask for a template if none is given).
- Content in simple markdown: '# ' title, '## ' headings, '- ' bullets, other lines are paragraphs.
- Both tools save a NEW file: end with a short message naming the exact file the tool returned and what changed."""


def document_editor(state: dict) -> dict:
    answer = run_tool_agent("document_editor", SYSTEM, state["task"], TOOLS)
    print(f"[document_editor] {answer[:80]}...")
    return {"agent_results": {"document_editor": answer}}


if __name__ == "__main__":
    import sys

    from docx.shared import Pt

    from userdata import use_temp_data, workspace_dir

    use_temp_data()  # test files go to a temp folder, never the real workspace
    sys.stdout.reconfigure(encoding="utf-8")
    ws = workspace_dir()

    # A template with a header (company), a footer, a body line and trailing blank lines
    t = Document()
    t.sections[0].header.paragraphs[0].text = "ACME Corp · 1 Main Street"
    t.sections[0].footer.paragraphs[0].text = "acme.example · Confidential"
    t.add_paragraph("ACME letter")
    t.add_paragraph("")
    t.add_paragraph("")
    t.save(ws / "acme_template.docx")

    # 1. fill_template keeps header/footer + body, drops the blank lines, writes content in the template
    msg = fill_template.invoke({"template_file": "acme_template.docx",
                                "content": "# Offer\n## Terms\n- Salary: **100**\nThanks."})
    out = Document(ws / "acme_template_filled.docx")
    assert "acme_template_filled.docx" in msg, msg
    assert out.sections[0].header.paragraphs[0].text == "ACME Corp · 1 Main Street"
    assert out.sections[0].footer.paragraphs[0].text == "acme.example · Confidential"
    texts = [p.text for p in out.paragraphs]
    assert texts == ["ACME letter", "Offer", "Terms", "Salary: 100", "Thanks."], texts
    assert out.paragraphs[1].style.name == "Heading 1" and out.paragraphs[3].runs[1].bold
    assert "acme_template_filled_2.docx" in fill_template.invoke({"template_file": "acme_template.docx",
                                                                  "content": "x"})  # never overwrites

    # 2. A template WITHOUT heading/bullet styles still works (bold heading, "• " bullets)
    bare = Document()
    for name in ("Heading 1", "List Bullet"):
        el = bare.styles[name].element
        el.getparent().remove(el)
    bare.save(ws / "bare.docx")
    fill_template.invoke({"template_file": "bare.docx", "content": "# Title\n- point"})
    bp = Document(ws / "bare_filled.docx").paragraphs
    assert bp[-2].text == "Title" and bp[-2].runs[0].bold and bp[-1].text == "• point", [p.text for p in bp]

    # 3. edit_document: formatting kept, runs spanning, tables and header, all-or-nothing
    d = Document()
    d.sections[0].header.paragraphs[0].text = "Draft 2025"
    p = d.add_paragraph()
    r = p.add_run("Dear John,")
    r.bold, r.font.size = True, Pt(14)
    p2 = d.add_paragraph()
    p2.add_run("Start date: 1 ")
    p2.add_run("May")  # "1 May" spans two runs
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Fee: 500"
    d.save(ws / "letter.docx")
    msg = edit_document.invoke({"filename": "letter.docx", "replacements": [
        {"find": "John", "replace": "Asha"}, {"find": "1 May", "replace": "3 June"},
        {"find": "500", "replace": "650"}, {"find": "2025", "replace": "2026"}]})
    e = Document(ws / "letter_edited.docx")
    assert "letter_edited.docx" in msg and "4 replacement" in msg, msg
    assert e.paragraphs[0].text == "Dear Asha," and e.paragraphs[0].runs[0].bold
    assert e.paragraphs[0].runs[0].font.size == Pt(14), "formatting must survive"
    assert e.paragraphs[1].text == "Start date: 3 June" and e.tables[0].cell(0, 0).text == "Fee: 650"
    assert e.sections[0].header.paragraphs[0].text == "Draft 2026"
    assert Document(ws / "letter.docx").paragraphs[0].text == "Dear John,", "the original must be untouched"
    try:
        edit_document.invoke({"filename": "letter.docx", "replacements": [
            {"find": "John", "replace": "X"}, {"find": "not there", "replace": "Y"}]})
        raise AssertionError("a missing find must fail")
    except ValueError:
        pass
    assert not (ws / "letter_edited_2.docx").exists(), "all-or-nothing: nothing saved"
    try:
        edit_document.invoke({"filename": "survey_report.pdf", "replacements": [{"find": "a", "replace": "b"}]})
        raise AssertionError("PDFs can't be edited in place")
    except FileNotFoundError:
        pass
    print("tool checks OK")
    if "--offline" in sys.argv:
        sys.exit()

    # 4. The agent: rewrite a PDF into the template
    before = {f.name for f in ws.iterdir()}
    ans = document_editor({"task": "Rewrite survey_report.pdf as a short formal summary in acme_template.docx."})
    print(ans["agent_results"]["document_editor"])
    made = [f.name for f in ws.iterdir() if f.name not in before and f.suffix == ".docx"]  # the agent may name it
    assert len(made) == 1, made
    text = read_document.invoke({"filename": made[0]})
    assert "82" in text and Document(ws / made[0]).sections[0].header.paragraphs[0].text.startswith("ACME"), text
    print("\nDocument Editor OK")
