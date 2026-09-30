"""Document Agent: reads, summarizes, creates and edits documents in the workspace/ folder.

    task -> tool-calling loop (list / read / create / append) -> "Created X" / summary -> result

Run the standalone test from the project root:  python -m agents.document
"""
from pathlib import Path

from docx import Document
from langchain_core.tools import tool
from pypdf import PdfReader

from agents.tool_agent import run_tool_agent
from userdata import workspace_dir

MAX_CHARS = 20_000      # read_document shows this much; longer files -> summarize_long_document (map-reduce)
CHUNK_CHARS = 10_000    # one "map" step's worth of text
MAX_CHUNKS = 15         # ~150k characters (~40 pages); beyond that the cost/time is too high for one request


def _safe_path(filename: str) -> Path:
    """SECURITY: the LLM picks file names, so only allow plain names inside workspace/."""
    workspace = workspace_dir()  # the CURRENT user's workspace
    path = (workspace / filename).resolve()
    if path.parent != workspace:  # blocks "../secrets.txt", "C:/Windows/...", sub-folders, other users' folders
        raise ValueError(f"'{filename}' must be a plain file name inside the workspace folder.")
    return path


def _add_text(paragraph, text: str) -> None:
    """'a **b** c' -> runs 'a ', bold 'b', ' c' (every odd piece after splitting on ** is bold)."""
    for i, piece in enumerate(text.split("**")):
        if piece:
            paragraph.add_run(piece).bold = i % 2 == 1


def _write_markdown(doc, content: str) -> None:
    """Turn simple markdown ('# ', '## ', '- ', **bold**) into Word headings, bullets and paragraphs."""
    for line in content.splitlines():
        line = line.strip()
        if line.startswith("## "):
            doc.add_heading(line[3:].replace("**", ""), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].replace("**", ""), level=1)
        elif line.startswith(("- ", "* ")):
            _add_text(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line:
            _add_text(doc.add_paragraph(), line)


# --- TOOLS: the docstring is what the LLM reads to decide when/how to call each one ---
@tool
def list_files() -> str:
    """List the files in the user's workspace folder."""
    # skip Office lock files ("~$deck.pptx" exists while the file is open in PowerPoint) and hidden files
    return "\n".join(p.name for p in workspace_dir().iterdir()
                     if p.is_file() and not p.name.startswith(("~$", "."))) or "(workspace is empty)"


@tool
def read_document(filename: str) -> str:
    """Read the text of a .docx, .pdf, .txt or .md file from the workspace (first 20,000 characters)."""
    text = _extract_text(filename)
    if len(text) <= MAX_CHARS:
        return text
    return (text[:MAX_CHARS] + f"\n[...truncated: showing {MAX_CHARS:,} of {len(text):,} characters. "
            "To cover the WHOLE file, use summarize_long_document.]")


def _extract_text(filename: str) -> str:
    path = _safe_path(filename)
    if not path.exists():
        raise FileNotFoundError(f"'{filename}' not found. Use list_files to see what exists.")
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return "\n".join(p.text for p in Document(path).paragraphs)
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8")
    raise ValueError(f"Unsupported file type '{suffix}'. Supported: .docx, .pdf, .txt, .md")


def _chunks(text: str, size: int = CHUNK_CHARS) -> list[str]:
    """Split at paragraph boundaries (never mid-sentence if avoidable), each chunk <= size characters."""
    chunks, current = [], ""
    for para in text.split("\n"):
        while len(para) > size:  # a single giant paragraph: hard-split it
            chunks.append(para[:size])
            para = para[size:]
        if len(current) + len(para) + 1 > size and current:
            chunks.append(current)
            current = ""
        current += para + "\n"
    return chunks + [current] if current.strip() else chunks


_summarizer = None


def summarizer():
    """The LLM used for map-reduce (created lazily, so tests can swap in a fake)."""
    global _summarizer
    if _summarizer is None:
        from llm import get_llm
        _summarizer = get_llm()
    return _summarizer


@tool
def summarize_long_document(filename: str, focus: str = "") -> str:
    """Summarize a document of ANY length (map-reduce: summarize each part, then combine).
    Use it when read_document says the file was truncated. `focus` = what to pay attention to (optional)."""
    from langchain_core.messages import HumanMessage, SystemMessage

    text = _extract_text(filename)
    parts = _chunks(text)
    if len(parts) > MAX_CHUNKS:
        raise ValueError(f"{filename} is too long ({len(text):,} characters, {len(parts)} parts; max {MAX_CHUNKS}).")
    want = f" Pay special attention to: {focus}." if focus else ""
    # MAP: each part on its own. ponytail: sequential to respect rate limits; could run a few in parallel.
    notes = []
    for i, part in enumerate(parts, 1):
        notes.append(summarizer().invoke([
            SystemMessage("Summarize this PART of a longer document in at most 8 bullet points. Keep every number, "
                          "name, date and decision exactly; don't add anything that isn't in the text." + want),
            HumanMessage(f"Part {i} of {len(parts)}:\n\n{part}"),
        ]).text)
    if len(notes) == 1:
        return f"Summary of {filename}:\n{notes[0]}"
    # REDUCE: combine the part-summaries (small enough now) into one summary
    combined = summarizer().invoke([
        SystemMessage("Combine these part-summaries of ONE document into a single structured summary. Keep all key "
                      "numbers, names, dates and decisions; remove repetition; don't add anything new." + want),
        HumanMessage("\n\n".join(f"[Part {i}]\n{n}" for i, n in enumerate(notes, 1))),
    ]).text
    return f"Summary of {filename} ({len(text):,} characters, {len(parts)} parts):\n{combined}"


@tool
def create_document(filename: str, content: str) -> str:
    """Create a NEW Word (.docx) file. `content` is simple markdown: '# ' title, '## ' headings,
    '- ' bullet points, other lines are paragraphs. Refuses to overwrite an existing file."""
    path = _safe_path(filename)
    if path.suffix.lower() != ".docx":
        raise ValueError("filename must end with .docx")
    if path.exists():  # overwriting needs human approval -> Phase 9
        raise FileExistsError(f"'{filename}' already exists. Pick a new name or use append_to_document.")
    doc = Document()
    _write_markdown(doc, content)
    doc.save(path)
    return f"Created {filename}"


@tool
def append_to_document(filename: str, content: str) -> str:
    """Add content (same simple markdown as create_document) to the END of an existing .docx file."""
    path = _safe_path(filename)
    if path.suffix.lower() != ".docx" or not path.exists():
        raise FileNotFoundError(f"'{filename}' is not an existing .docx file in the workspace.")
    doc = Document(path)
    _write_markdown(doc, content)
    doc.save(path)
    return f"Appended to {filename}"


TOOLS = [list_files, read_document, summarize_long_document, create_document, append_to_document]

SYSTEM = """You are the Document Agent. You read, summarize, create and edit documents in the user's workspace.
- If unsure of a file name, call list_files first.
- To summarize, call read_document, then write the summary yourself. Only use facts from the document.
- If read_document says the file was truncated, use summarize_long_document instead, so the WHOLE file is covered.
- Only create or change files when the task asks for a file. To read, extract or summarize, reply with the text.
- Write document content in simple markdown: '# ' title, '## ' headings, '- ' bullets.
- End with a short message: what you did and which file(s) you created or changed."""


def document_agent(state: dict) -> dict:
    answer = run_tool_agent("document_agent", SYSTEM, state["task"], TOOLS)
    print(f"[document_agent] {answer[:80]}...")
    return {"agent_results": {"document_agent": answer}}


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding="utf-8")

    # 1. Tools alone (no LLM): security checks must hold no matter what the LLM asks for
    for bad in ["../secrets.txt", "C:/Windows/win.ini", "sub/x.docx"]:
        try:
            _safe_path(bad)
            raise AssertionError(f"path not blocked: {bad}")
        except ValueError:
            pass
    print("path checks OK")

    from userdata import use_temp_data
    use_temp_data()  # test files go to a temp folder, never the real workspace
    guide = workspace_dir() / "welcome_guide.docx"
    guide.unlink(missing_ok=True)  # make the test repeatable

    # 2. Read + summarize a PDF
    ans = document_agent({"task": "Summarize survey_report.pdf in 3 bullet points."})["agent_results"]["document_agent"]
    assert "82" in ans, ans

    # 3. Create a Word document
    document_agent({"task": "Create welcome_guide.docx: a short welcome guide for new employees with the "
                            "sections 'Before day one', 'First week' and 'Who to ask'."})
    text = read_document.invoke({"filename": "welcome_guide.docx"})
    assert guide.exists() and "first week" in text.lower() and "**" not in text, text

    # 4. Edit it (append a section)
    document_agent({"task": "Add a section called 'Useful links' with 2 bullet points to welcome_guide.docx."})
    assert "useful links" in read_document.invoke({"filename": "welcome_guide.docx"}).lower()

    print("\nDocument agent OK")
