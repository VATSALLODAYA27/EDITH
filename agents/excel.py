"""Excel Agent: reads, analyzes, creates and edits .xlsx spreadsheets (formulas + charts) in workspace/.

    task -> tool-calling loop (read / analyze / write / update / chart) -> answer with real numbers -> result

KEY IDEA: LLMs are bad at arithmetic. The LLM decides WHAT to compute; analyze_sheet does the maths in Python.

Run the standalone test from the project root:  python -m agents.excel
"""
from collections import defaultdict

from langchain_core.tools import tool
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from agents.document import _safe_path, list_files  # same workspace + same security check
from agents.tool_agent import run_tool_agent

CHARTS = {"bar": BarChart, "line": LineChart, "pie": PieChart}


# --- helpers ---
def _open(filename: str):
    path = _safe_path(filename)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx files are supported.")
    if not path.exists():
        raise FileNotFoundError(f"'{filename}' not found. Use list_files to see what exists.")
    return path, load_workbook(path)  # note: formulas load as text like '=SUM(B2:B5)', never as results


def _sheet(wb, sheet_name: str):
    if not sheet_name:
        return wb.active
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"No sheet '{sheet_name}'. Sheets: {wb.sheetnames}")
    return wb[sheet_name]


def _column(ws, header: str) -> int:
    """Header name (row 1, case-insensitive) -> column number."""
    headers = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
    if header.strip().lower() not in headers:
        raise ValueError(f"No column '{header}'. Columns: {[c.value for c in ws[1]]}")
    return headers.index(header.strip().lower()) + 1


def _number(value):
    """LLMs often send numbers as text ('1,200'). Store them as real numbers so Excel can do maths on them."""
    if isinstance(value, str) and not value.startswith("="):
        try:
            f = float(value.replace(",", ""))
            return int(f) if f.is_integer() else f
        except ValueError:
            pass
    return value


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _fmt(x: float) -> str:
    return f"{x:,.0f}" if float(x).is_integer() else f"{x:,.2f}"


# --- TOOLS ---
@tool
def read_sheet(filename: str, sheet_name: str = "", max_rows: int = 50) -> str:
    """Show a sheet's contents with row numbers and column letters. Empty sheet_name = first sheet."""
    _, wb = _open(filename)
    ws = _sheet(wb, sheet_name)
    letters = ", ".join(f"{get_column_letter(c.column)}={c.value}" for c in ws[1])
    lines = [f"Sheets: {wb.sheetnames}. Sheet '{ws.title}': {ws.max_row} rows. Columns: {letters}"]
    for i, row in enumerate(ws.iter_rows(max_row=max_rows, values_only=True), start=1):
        lines.append(f"{i}: " + " | ".join("" if v is None else str(v) for v in row))
    if ws.max_row > max_rows:
        lines.append(f"... {ws.max_row - max_rows} more rows. Use analyze_sheet for totals/averages.")
    return "\n".join(lines)


@tool
def analyze_sheet(filename: str, value_column: str, group_by: str = "", sheet_name: str = "") -> str:
    """Compute sum, average, min, max and count of a numeric column, optionally per group
    (e.g. value_column='Revenue', group_by='Region'). ALWAYS use this instead of doing maths yourself."""
    _, wb = _open(filename)
    ws = _sheet(wb, sheet_name)
    v_col = _column(ws, value_column)
    g_col = _column(ws, group_by) if group_by else None

    groups, skipped = defaultdict(list), 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        value = row[v_col - 1]
        if not _is_num(value):  # formulas (e.g. a Total row), blanks, text -> not counted
            skipped += value is not None
            continue
        groups[row[g_col - 1] if g_col else "ALL"].append(value)

    lines = [f"{value_column}" + (f" by {group_by}" if group_by else "") + " (sorted by sum):"]
    for key, vals in sorted(groups.items(), key=lambda kv: -sum(kv[1])):
        lines.append(f"- {key}: sum={_fmt(sum(vals))}, avg={_fmt(sum(vals) / len(vals))}, "
                     f"min={_fmt(min(vals))}, max={_fmt(max(vals))}, count={len(vals)}")
    if skipped:
        lines.append(f"(skipped {skipped} non-numeric cells, e.g. formulas or text)")
    return "\n".join(lines)


@tool
def write_sheet(filename: str, sheet_name: str, rows: list[list[str | float]]) -> str:
    """Write rows (first row = headers) to a NEW sheet. Creates the .xlsx file if it doesn't exist.
    Values starting with '=' are Excel formulas, e.g. '=SUM(B2:B5)'. Refuses to replace an existing sheet."""
    path = _safe_path(filename)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("filename must end with .xlsx")
    if path.exists():
        wb = load_workbook(path)
        if sheet_name in wb.sheetnames:  # replacing data needs human approval -> Phase 9
            raise ValueError(f"Sheet '{sheet_name}' already exists. Use update_cells or another sheet name.")
        ws = wb.create_sheet(sheet_name)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
    for row in rows:
        ws.append([_number(v) for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(path)
    return f"Wrote {len(rows)} rows to sheet '{sheet_name}' in {filename}"


@tool
def update_cells(filename: str, updates: dict[str, str | float], sheet_name: str = "") -> str:
    """Set specific cells in an existing sheet, e.g. {"B6": "=SUM(B2:B5)", "A6": "Total"}.
    Values starting with '=' are Excel formulas."""
    path, wb = _open(filename)
    ws = _sheet(wb, sheet_name)
    for ref, value in updates.items():
        coordinate_from_string(ref)  # raises on invalid references like "Z" or "hello"
        ws[ref] = _number(value)
    wb.save(path)
    return f"Updated {len(updates)} cell(s) in '{ws.title}' of {filename}"


@tool
def add_chart(filename: str, category_column: str, value_column: str, chart_type: str = "bar",
              title: str = "", sheet_name: str = "") -> str:
    """Add a chart (bar, line or pie) of value_column per category_column, using the sheet's numeric rows.
    For aggregated charts (e.g. revenue by region) first write a summary sheet, then chart that sheet."""
    if chart_type not in CHARTS:
        raise ValueError(f"chart_type must be one of {list(CHARTS)}")
    path, wb = _open(filename)
    ws = _sheet(wb, sheet_name)
    c_col, v_col = _column(ws, category_column), _column(ws, value_column)
    numeric_rows = [r for r in range(2, ws.max_row + 1) if _is_num(ws.cell(r, v_col).value)]
    if not numeric_rows:
        raise ValueError(f"Column '{value_column}' has no numeric values to chart.")
    last = numeric_rows[-1]  # stops before a formula Total row, which would dwarf every other bar

    chart = CHARTS[chart_type]()
    chart.title = title or f"{value_column} by {category_column}"
    chart.add_data(Reference(ws, min_col=v_col, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=c_col, min_row=2, max_row=last))
    anchor = f"{get_column_letter(ws.max_column + 2)}{2 + 16 * len(ws._charts)}"  # right of data, stacked
    ws.add_chart(chart, anchor)
    wb.save(path)
    return f"Added {chart_type} chart '{chart.title}' at {anchor} in sheet '{ws.title}' of {filename}"


TOOLS = [list_files, read_sheet, analyze_sheet, write_sheet, update_cells, add_chart]

SYSTEM = """You are the Excel Agent. You work with .xlsx spreadsheets in the user's workspace.
- Look before you act: call read_sheet to see the headers and layout first.
- NEVER do arithmetic yourself. Use analyze_sheet for totals, averages, min/max, and quote its numbers exactly.
- Don't add units or currency symbols that aren't in the data.
- In write_sheet/update_cells, values starting with '=' are Excel formulas (e.g. '=SUM(B2:B5)').
- To chart aggregated data (e.g. revenue by region): analyze_sheet -> write_sheet a small summary sheet -> add_chart on it.
- End with a short message: what you did, the key numbers, and which files/sheets you changed."""


def excel_agent(state: dict) -> dict:
    answer = run_tool_agent("excel_agent", SYSTEM, state["task"], TOOLS)
    print(f"[excel_agent] {answer[:80]}...")
    return {"agent_results": {"excel_agent": answer}}


if __name__ == "__main__":
    import sys
    import zipfile

    from userdata import use_temp_data, workspace_dir

    use_temp_data()
    WORKSPACE = workspace_dir()

    sys.stdout.reconfigure(encoding="utf-8")

    # Ground truth computed independently of our tools, to check the agent's numbers
    sales = load_workbook(WORKSPACE / "sales.xlsx").active
    totals = defaultdict(int)
    for _, region, _, _, revenue in sales.iter_rows(min_row=2, values_only=True):
        totals[region] += revenue
    best = max(totals, key=totals.get)

    # 1. Tools alone (no LLM)
    out = analyze_sheet.invoke({"filename": "sales.xlsx", "value_column": "Revenue", "group_by": "Region"})
    print(out)
    assert out.splitlines()[1].startswith(f"- {best}: sum={_fmt(totals[best])}")

    (WORKSPACE / "t.xlsx").unlink(missing_ok=True)
    write_sheet.invoke({"filename": "t.xlsx", "sheet_name": "S", "rows": [["n"], ["1,500"], ["2"], ["=SUM(A2:A3)"]]})
    assert "=SUM(A2:A3)" in read_sheet.invoke({"filename": "t.xlsx"})  # formula stored, not calculated
    assert "sum=1,502" in analyze_sheet.invoke({"filename": "t.xlsx", "value_column": "n"})  # '1,500' -> number
    (WORKSPACE / "t.xlsx").unlink()
    print("tool checks OK\n")

    # 2. Analysis question -> must use analyze_sheet and report the right number
    ans = excel_agent({"task": "Which region had the highest total revenue in sales.xlsx, and how much?"})
    ans = ans["agent_results"]["excel_agent"]
    assert best in ans and str(totals[best]) in ans.replace(",", ""), ans

    # 3. Create a workbook with a formula
    budget = WORKSPACE / "budget.xlsx"
    budget.unlink(missing_ok=True)
    excel_agent({"task": "Create budget.xlsx with a sheet 'Budget': columns Item and Cost for laptop 80000, "
                         "chair 12000, desk 18000, monitor 15000, then a Total row using a SUM formula."})
    cells = [c.value for row in load_workbook(budget)["Budget"].iter_rows() for c in row]
    assert any(isinstance(v, str) and v.upper().startswith("=SUM(") for v in cells), cells

    # 4. Add a chart
    excel_agent({"task": "Add a bar chart of Cost by Item to budget.xlsx."})
    assert any("charts/chart" in n for n in zipfile.ZipFile(budget).namelist()), "no chart in file"

    # 5. Multi-step: an aggregated chart needs analyze -> summary sheet -> chart (on a copy, keep the sample clean)
    import shutil
    copy = WORKSPACE / "sales_copy.xlsx"
    shutil.copy(WORKSPACE / "sales.xlsx", copy)
    excel_agent({"task": "Add a bar chart of total revenue per region to sales_copy.xlsx."})
    wb = load_workbook(copy)
    assert len(wb.sheetnames) == 2, f"expected a summary sheet, got {wb.sheetnames}"
    summary = {r[0]: r[1] for r in wb[wb.sheetnames[1]].iter_rows(min_row=2, values_only=True) if r[0]}
    assert all(summary.get(k) == v for k, v in totals.items()), summary  # the chart shows the RIGHT totals
    assert any("charts/chart" in n for n in zipfile.ZipFile(copy).namelist())
    copy.unlink()

    print("\nExcel agent OK")
