"""PPT Agent: creates and edits PowerPoint (.pptx) presentations in workspace/.

    task -> tool-calling loop (read / create / add / update / delete slides) -> "Created X" -> result

KEY IDEA: the LLM writes CONTENT as structured data (title, bullets, notes per slide);
Python renders it into real layouts and ENFORCES limits (e.g. max bullets) so slides stay readable.

Run the standalone test from the project root:  python -m agents.ppt
"""
from langchain_core.tools import tool
from pptx import Presentation
from pydantic import BaseModel, Field

from agents.document import _safe_path, list_files  # same workspace + same security check
from agents.tool_agent import run_tool_agent

TITLE_LAYOUT, CONTENT_LAYOUT = 0, 1  # layouts in python-pptx's default template
MAX_BULLETS, MAX_BULLET_CHARS = 6, 120  # more than this overflows a default slide


class Slide(BaseModel):
    title: str
    bullets: list[str] = Field(default_factory=list, description=f"At most {MAX_BULLETS} short bullet points.")
    notes: str = Field(default="", description="Optional speaker notes.")


# --- helpers ---
def _open(filename: str):
    path = _safe_path(filename)
    if path.suffix.lower() != ".pptx":
        raise ValueError("Only .pptx files are supported.")
    if not path.exists():
        raise FileNotFoundError(f"'{filename}' not found. Use list_files to see what exists.")
    return path, Presentation(path)


def _check(slide: Slide) -> None:
    """Tool-enforced limits: the error goes back to the LLM, which then splits the content."""
    if len(slide.bullets) > MAX_BULLETS:
        raise ValueError(f"Slide '{slide.title}' has {len(slide.bullets)} bullets; max is {MAX_BULLETS}. "
                         "Split it into more slides.")
    too_long = [b for b in slide.bullets if len(b) > MAX_BULLET_CHARS]
    if too_long:
        raise ValueError(f"Bullets must be under {MAX_BULLET_CHARS} characters. Shorten: {too_long[0][:60]}...")


def _slide_index(prs, slide_number: int) -> int:
    """1-based slide number (what humans and the LLM use) -> 0-based index, with a helpful error."""
    if not 1 <= slide_number <= len(prs.slides):
        raise ValueError(f"slide_number must be 1..{len(prs.slides)}")
    return slide_number - 1


def _fill(slide, spec: Slide) -> None:
    slide.shapes.title.text = spec.title
    body = slide.placeholders[1].text_frame  # layout 1: content box; layout 0: subtitle
    body.text = spec.bullets[0] if spec.bullets else ""
    for bullet in spec.bullets[1:]:
        body.add_paragraph().text = bullet
    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes


def _add(prs, spec: Slide, layout: int = CONTENT_LAYOUT):
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    _fill(slide, spec)
    return slide


def _move_last_to(prs, position: int) -> None:
    """python-pptx has no 'move slide' API; slide order is the order of <p:sldId> elements in the XML."""
    order = prs.slides._sldIdLst
    last = order[-1]
    order.remove(last)
    order.insert(position, last)


# --- TOOLS ---
@tool
def read_presentation(filename: str) -> str:
    """Show every slide: its number, title, text and speaker notes."""
    _, prs = _open(filename)
    lines = [f"{filename}: {len(prs.slides)} slides"]
    for n, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else "(no title)"
        text = [p.text for s in slide.shapes if s.has_text_frame and s != slide.shapes.title
                for p in s.text_frame.paragraphs if p.text]
        lines.append(f"Slide {n}: {title}" + "".join(f"\n  - {t}" for t in text))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
            lines.append(f"  notes: {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(lines)


@tool
def create_presentation(filename: str, title: str, slides: list[Slide], subtitle: str = "") -> str:
    """Create a NEW .pptx: a title slide (title + subtitle) followed by one content slide per item in `slides`.
    Refuses to overwrite an existing file."""
    path = _safe_path(filename)
    if path.suffix.lower() != ".pptx":
        raise ValueError("filename must end with .pptx")
    if path.exists():  # overwriting needs human approval -> Phase 9
        raise FileExistsError(f"'{filename}' already exists. Pick a new name or edit it with the other tools.")
    for spec in slides:
        _check(spec)
    prs = Presentation()
    _add(prs, Slide(title=title, bullets=[subtitle] if subtitle else []), layout=TITLE_LAYOUT)
    for spec in slides:
        _add(prs, spec)
    prs.save(path)
    return f"Created {filename} with {len(prs.slides)} slides"


@tool
def add_slide(filename: str, slide: Slide, position: int = 0) -> str:
    """Add a content slide. position = slide number it should become (1 = first); 0 = at the end."""
    path, prs = _open(filename)
    _check(slide)
    _add(prs, slide)
    if position:
        _move_last_to(prs, _slide_index(prs, position))
    prs.save(path)
    return f"Added slide '{slide.title}' to {filename} ({len(prs.slides)} slides now)"


@tool
def update_slide(filename: str, slide_number: int, title: str = "", bullets: list[str] | None = None,
                 notes: str | None = None) -> str:
    """Change an existing slide (slide_number starts at 1). Only the fields you pass change; the rest keep
    their current content (e.g. pass only `title` to rename a slide)."""
    path, prs = _open(filename)
    target = prs.slides[_slide_index(prs, slide_number)]
    body = target.placeholders[1].text_frame if len(target.placeholders) > 1 else None
    current = [p.text for p in body.paragraphs if p.text] if body else []
    spec = Slide(title=title or target.shapes.title.text, bullets=current if bullets is None else bullets)
    _check(spec)
    for shape in target.placeholders:  # clear old text, keep the layout
        if shape.has_text_frame:
            shape.text_frame.clear()
    _fill(target, spec)
    if notes is not None:  # None = keep existing notes; "" = clear them
        target.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return f"Updated slide {slide_number} of {filename}"


@tool
def delete_slide(filename: str, slide_number: int) -> str:
    """Request deleting one slide (slide_number starts at 1). Deleting can't be undone, so the slide is only
    removed after the user approves."""
    from approvals import queue_file_action  # imported here to avoid a circular import

    _, prs = _open(filename)
    title = prs.slides[_slide_index(prs, slide_number)].shapes.title
    title = title.text if title is not None else ""
    action_id = queue_file_action({"filename": filename, "slide_number": slide_number, "title": title})
    return f"Deletion of slide {slide_number} ('{title}') requested as file:{action_id}. NOT deleted until the user approves."


def delete_slide_now(filename: str, slide_number: int, title: str | None = None) -> str:
    """HUMAN-APPROVED deletion (called by approvals.execute, never by the LLM).
    Checks the title too: if an earlier deletion shifted the numbers, find the slide by title instead."""
    path, prs = _open(filename)
    titles = [s.shapes.title.text if s.shapes.title is not None else "" for s in prs.slides]
    index = _slide_index(prs, slide_number)
    if title is not None and titles[index] != title:
        if title not in titles:
            raise ValueError(f"Slide '{title}' is no longer in {filename}; nothing deleted.")
        index = titles.index(title)
    order = prs.slides._sldIdLst
    entry = order[index]
    prs.part.drop_rel(entry.rId)  # unlink the slide part...
    order.remove(entry)           # ...and remove it from the slide order
    prs.save(path)
    return f"Deleted slide '{titles[index]}' from {filename} ({len(prs.slides)} slides left)"


TOOLS = [list_files, read_presentation, create_presentation, add_slide, update_slide, delete_slide]

SYSTEM = f"""You are the PPT Agent. You create and edit PowerPoint (.pptx) presentations in the user's workspace.
- Before editing an existing file, call read_presentation to see its slides and their numbers.
- Good slides: a short title, at most {MAX_BULLETS} concise bullets, one idea per slide. Use speaker notes for detail.
- Use only facts given in the task; don't invent numbers, names or dates.
- Slide numbers change after adding/deleting slides: re-read the presentation if you need them again.
- delete_slide only REQUESTS a deletion; the user approves it afterwards. Report it as requested, not as done.
- End with a short message: what you did and which file you created or changed."""


def ppt_agent(state: dict) -> dict:
    answer = run_tool_agent("ppt_agent", SYSTEM, state["task"], TOOLS)
    print(f"[ppt_agent] {answer[:80]}...")
    return {"agent_results": {"ppt_agent": answer}}


def _titles(filename: str) -> list[str]:
    _, prs = _open(filename)
    return [s.shapes.title.text for s in prs.slides]


if __name__ == "__main__":
    import sys

    from userdata import use_temp_data, workspace_dir

    use_temp_data()
    WORKSPACE = workspace_dir()

    sys.stdout.reconfigure(encoding="utf-8")

    # 1. Tools alone (no LLM)
    (WORKSPACE / "t.pptx").unlink(missing_ok=True)
    create_presentation.invoke({"filename": "t.pptx", "title": "T", "subtitle": "s",
                                "slides": [{"title": "A", "bullets": ["a1"]}, {"title": "C"}]})
    add_slide.invoke({"filename": "t.pptx", "slide": {"title": "B", "notes": "note"}, "position": 3})
    assert _titles("t.pptx") == ["T", "A", "B", "C"], _titles("t.pptx")
    update_slide.invoke({"filename": "t.pptx", "slide_number": 2, "title": "A2"})  # rename only
    assert "- a1" in read_presentation.invoke({"filename": "t.pptx"}), "renaming must keep the bullets"
    update_slide.invoke({"filename": "t.pptx", "slide_number": 2, "bullets": ["x"]})  # bullets only
    msg = delete_slide.invoke({"filename": "t.pptx", "slide_number": 4})  # only a REQUEST now
    assert _titles("t.pptx") == ["T", "A2", "B", "C"] and "NOT deleted" in msg, msg
    import approvals
    approvals.execute(f"file:{msg.split('file:')[1].split('.')[0]}")  # what the approval node does on "approve"
    assert _titles("t.pptx") == ["T", "A2", "B"], _titles("t.pptx")
    deck = read_presentation.invoke({"filename": "t.pptx"})
    assert "notes: note" in deck and "- x" in deck and "- a1" not in deck, deck
    try:
        _check(Slide(title="x", bullets=["b"] * 7))
        raise AssertionError("7 bullets should be rejected")
    except ValueError:
        pass
    (WORKSPACE / "t.pptx").unlink()
    print("tool checks OK\n")

    # 2. Create a deck from facts given in the task
    deck = WORKSPACE / "onboarding.pptx"
    deck.unlink(missing_ok=True)
    ppt_agent({"task": "Create onboarding.pptx: title slide 'Welcome to Nimbus Labs', then one slide each on "
                       "Vacation (24 paid days a year, 5 carry over), Remote work (up to 3 days a week, "
                       "Tue/Thu in office) and Expenses (submit receipts within 30 days)."})
    titles = _titles("onboarding.pptx")
    print(titles)
    assert len(titles) == 4 and any("vacation" in t.lower() for t in titles), titles

    # 3. Edit: add a slide at the end + rename slide 2
    ppt_agent({"task": "In onboarding.pptx, add a final slide titled 'Questions?' and rename slide 2 to 'Time off'."})
    titles = _titles("onboarding.pptx")
    print(titles)
    assert titles[1] == "Time off" and titles[-1] == "Questions?" and len(titles) == 5, titles

    # 4. Delete by meaning (agent must read first to find the right slide) -> only QUEUED until approved
    ppt_agent({"task": "Remove the slide about expenses from onboarding.pptx."})
    queued = [a for a in approvals.list_pending() if a["kind"] == "file"]
    print(queued)
    assert len(_titles("onboarding.pptx")) == 5 and len(queued) == 1 and "expense" in queued[0]["summary"].lower()
    approvals.execute(queued[0]["id"])
    titles = _titles("onboarding.pptx")
    assert len(titles) == 4 and not any("expense" in t.lower() for t in titles), titles

    # 5. Limits: 9 tips "on one slide" -> tool rejects >6 bullets -> agent must split across slides
    (WORKSPACE / "tips.pptx").unlink(missing_ok=True)
    tips = [f"Tip {i}" for i in range(1, 10)]
    ppt_agent({"task": f"Create tips.pptx (title 'Productivity tips') with one slide listing these tips: {tips}"})
    _, prs = _open("tips.pptx")
    bullets = [len(s.placeholders[1].text_frame.paragraphs) for s in list(prs.slides)[1:]]
    print("bullets per content slide:", bullets)
    assert max(bullets) <= MAX_BULLETS and sum(bullets) == 9, bullets
    (WORKSPACE / "tips.pptx").unlink()

    print("\nPPT agent OK")
